import argparse
from pathlib import Path
from html.parser import HTMLParser
from docx import Document
from docx.shared import Pt

parser = argparse.ArgumentParser()
parser.add_argument('--html', required=True)
parser.add_argument('--out', required=True)
args = parser.parse_args()
class TextParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.items = []
        self.current = None
    def handle_starttag(self, tag, attrs):
        if tag in ('h1', 'h2', 'h3', 'p', 'li'):
            self.current = [tag, []]
    def handle_data(self, data):
        if self.current:
            self.current[1].append(data)
    def handle_endtag(self, tag):
        if self.current and self.current[0] == tag:
            self.items.append((tag, ' '.join(''.join(self.current[1]).split())))
            self.current = None
parser_html = TextParser()
parser_html.feed(Path(args.html).read_text(encoding='utf-8'))
doc = Document()
style = doc.styles['Normal']
style.font.name = 'Arial'
style.font.size = Pt(10)
for tag, text in parser_html.items:
    if not text:
        continue
    if tag == 'h1':
        doc.add_heading(text, 0)
    elif tag == 'h2':
        doc.add_heading(text, 1)
    elif tag == 'h3':
        doc.add_heading(text, 2)
    elif tag == 'li':
        doc.add_paragraph(text, style='List Bullet')
    else:
        doc.add_paragraph(text)
doc.save(args.out)
